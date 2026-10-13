"""Document parser — public interface: parse(path) -> Document.

Detects format by magic bytes (not extension), then dispatches to
format-specific parsers. Enforces file-size, page, and decompression limits.
"""

from __future__ import annotations

import hashlib
import uuid
import zipfile
from pathlib import Path
from typing import Any

from tprm.common.config import get_settings
from tprm.common.errors import (
    FileTooLarge,
    IngestionError,
    TooManyPages,
    UnsupportedFormat,
    ZipBombDetected,
)
from tprm.common.logging import get_logger
from tprm.common.models import Block, BoundingBox, Document, DocumentMeta
from tprm.ingestion.ocr import ocr_image_bytes, ocr_pixmap
from tprm.ingestion.segmenter import segment

log = get_logger(__name__)

# ── Magic byte signatures ─────────────────────────────────────────────────────

_PDF_MAGIC = b"%PDF"
_DOCX_MAGIC = b"PK\x03\x04"  # ZIP-based; further inspected for content types
_XLSX_MAGIC = b"PK\x03\x04"  # same as DOCX

_OCR_SPARSE_THRESHOLD = 100  # chars per page


def parse(path: Path, vendor_id: str = "unknown") -> Document:
    """Parse a vendor document file and return a ``Document``.

    Args:
        path: Filesystem path to the file.
        vendor_id: The vendor this document belongs to.

    Returns:
        A ``Document`` with populated ``blocks`` and ``meta``.

    Raises:
        UnsupportedFormat: If the file type is not PDF, DOCX, or XLSX.
        FileTooLarge: If the file exceeds the configured size limit.
        ZipBombDetected: If decompression ratio is excessive.
        IngestionError: On any parsing failure.
    """
    settings = get_settings()
    policy = settings.policy

    raw = path.read_bytes()
    file_size = len(raw)

    max_size = policy.get("limits", {}).get("max_file_size_bytes", 52428800)
    if file_size > max_size:
        raise FileTooLarge(
            f"File {path.name!r} is {file_size} bytes, exceeds limit of {max_size}"
        )

    sha256 = hashlib.sha256(raw).hexdigest()
    doc_id = str(uuid.uuid4())

    fmt = _detect_format(raw, path)
    log.info("parse_start", doc_id=doc_id, filename=path.name, format=fmt, size_bytes=file_size)

    try:
        if fmt == "pdf":
            blocks, meta = _parse_pdf(raw, doc_id, policy)
        elif fmt == "docx":
            blocks, meta = _parse_docx(raw, doc_id, policy)
        elif fmt == "xlsx":
            blocks, meta = _parse_xlsx(raw, doc_id, policy)
        else:
            raise UnsupportedFormat(f"Unsupported format: {fmt!r}")
    except (UnsupportedFormat, FileTooLarge, TooManyPages, ZipBombDetected):
        raise
    except Exception as exc:
        raise IngestionError(f"Failed to parse {fmt!r} document: {exc}") from exc

    meta_obj = DocumentMeta(
        file_size_bytes=file_size,
        page_count=meta.get("page_count", 0),
        doc_format=fmt,  # type: ignore[arg-type]
        has_ocr_pages=meta.get("has_ocr_pages", False),
        has_hidden_sheets=meta.get("has_hidden_sheets", False),
    )

    doc = Document(
        doc_id=doc_id,
        vendor_id=vendor_id,
        filename=path.name,
        sha256=sha256,
        blocks=blocks,
        meta=meta_obj,
    )

    doc = segment(doc)
    log.info("parse_complete", doc_id=doc_id, block_count=len(doc.blocks))
    return doc


def _detect_format(raw: bytes, path: Path) -> str:
    """Detect format by magic bytes, NOT file extension."""
    if raw[:4] == _PDF_MAGIC:
        return "pdf"
    if raw[:4] == _DOCX_MAGIC:
        # Both DOCX and XLSX are ZIP. Distinguish by inspecting the ZIP entries.
        return _detect_zip_subtype(raw)
    raise UnsupportedFormat(
        f"Unrecognised file format (magic bytes: {raw[:4]!r}). "
        "Only PDF, DOCX, and XLSX are supported."
    )


def _detect_zip_subtype(raw: bytes) -> str:
    """Distinguish DOCX from XLSX by ZIP entry names."""
    try:
        import io
        with zipfile.ZipFile(io.BytesIO(raw)) as zf:
            names = zf.namelist()
    except zipfile.BadZipFile as exc:
        raise UnsupportedFormat("ZIP file is corrupt or not a valid DOCX/XLSX") from exc

    if any(n.startswith("xl/") or n == "[Content_Types].xml" and "xl/" in str(names) for n in names):
        # Check for Excel-specific entries
        if any(n.startswith("xl/worksheets/") for n in names):
            return "xlsx"
    if any(n.startswith("word/") for n in names):
        return "docx"
    if any(n.startswith("xl/") for n in names):
        return "xlsx"
    raise UnsupportedFormat("ZIP archive is neither a DOCX nor XLSX file")


def _check_zip_bomb(raw: bytes, path_name: str, ratio_limit: int = 50) -> None:
    """Raise ZipBombDetected if decompression ratio is excessive."""
    import io

    try:
        compressed_size = len(raw)
        total_uncompressed = 0
        with zipfile.ZipFile(io.BytesIO(raw)) as zf:
            for info in zf.infolist():
                total_uncompressed += info.file_size
                if compressed_size > 0 and total_uncompressed / compressed_size > ratio_limit:
                    raise ZipBombDetected(
                        f"Decompression ratio exceeds {ratio_limit}x for {path_name!r}"
                    )
    except ZipBombDetected:
        raise
    except Exception:
        pass  # Not a ZIP; non-applicable


# ── PDF parser ────────────────────────────────────────────────────────────────


def _parse_pdf(raw: bytes, doc_id: str, policy: dict[str, Any]) -> tuple[list[Block], dict[str, Any]]:
    """Parse a PDF using PyMuPDF. OCR fallback for sparse-text pages."""
    try:
        import fitz  # type: ignore[import-untyped]
    except ImportError as exc:
        raise IngestionError("PyMuPDF (fitz) is not installed") from exc

    max_pages = policy.get("limits", {}).get("max_pages", 500)
    sparse_threshold = policy.get("limits", {}).get("ocr_sparse_text_threshold", 100)
    conf_threshold = policy.get("limits", {}).get("ocr_confidence_threshold", 40)

    import io

    doc = fitz.open(stream=io.BytesIO(raw), filetype="pdf")
    page_count = doc.page_count

    if page_count > max_pages:
        raise TooManyPages(f"PDF has {page_count} pages, exceeds limit of {max_pages}")

    blocks: list[Block] = []
    has_ocr_pages = False

    for page_num in range(page_count):
        page = doc[page_num]
        native_text = page.get_text("text")

        if len(native_text.strip()) < sparse_threshold:
            # Sparse page → OCR
            has_ocr_pages = True
            pixmap = page.get_pixmap(dpi=300)
            ocr_text, mean_conf = ocr_pixmap(pixmap, doc_id, page_num, conf_threshold)
            if ocr_text.strip():
                blocks.append(
                    Block(
                        doc_id=doc_id,
                        page=page_num,
                        text=ocr_text,
                        source="ocr",
                        ocr_conf=mean_conf,
                    )
                )
            elif native_text.strip():
                blocks.extend(_native_pdf_blocks(page, doc_id, page_num))
        else:
            blocks.extend(_native_pdf_blocks(page, doc_id, page_num))

    doc.close()
    return blocks, {"page_count": page_count, "has_ocr_pages": has_ocr_pages}


def _native_pdf_blocks(page: Any, doc_id: str, page_num: int) -> list[Block]:
    """Extract native PDF text blocks with a whole-page fallback."""
    blocks: list[Block] = []
    raw_blocks = page.get_text("blocks")
    for b in raw_blocks:
        # b = (x0, y0, x1, y1, text, block_no, block_type)
        if len(b) < 5:
            continue
        x0, y0, x1, y1, text = b[0], b[1], b[2], b[3], b[4]
        text = str(text).strip()
        if not text:
            continue
        blocks.append(
            Block(
                doc_id=doc_id,
                page=page_num,
                text=text,
                bbox=BoundingBox(x0=x0, y0=y0, x1=x1, y1=y1),
                source="native",
            )
        )

    if not blocks:
        text = page.get_text("text").strip()
        if text:
            blocks.append(
                Block(
                    doc_id=doc_id,
                    page=page_num,
                    text=text,
                    source="native",
                )
            )
    return blocks


# ── DOCX parser ───────────────────────────────────────────────────────────────


def _parse_docx(raw: bytes, doc_id: str, policy: dict[str, Any]) -> tuple[list[Block], dict[str, Any]]:
    """Parse a DOCX using python-docx."""
    try:
        import io

        import docx  # type: ignore[import-untyped]
    except ImportError as exc:
        raise IngestionError("python-docx is not installed") from exc

    _check_zip_bomb(raw, "docx", policy.get("limits", {}).get("zip_bomb_ratio", 50))

    document = docx.Document(io.BytesIO(raw))
    blocks: list[Block] = []
    page_num = 0  # DOCX doesn't have real page numbers; use 0 throughout

    # Paragraphs
    for para in document.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        section_path: list[str] = []
        style_name = para.style.name if para.style else ""
        if _DOCX_HEADING_STYLES_RE.match(style_name):
            section_path = [text]
        blocks.append(
            Block(
                doc_id=doc_id,
                page=page_num,
                text=text,
                source="native",
                section_path=section_path,
            )
        )

    # Tables (cell-wise blocks)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                cell_text = cell.text.strip()
                if not cell_text:
                    continue
                blocks.append(
                    Block(
                        doc_id=doc_id,
                        page=page_num,
                        text=cell_text,
                        source="table_cell",
                    )
                )

    # Embedded images → OCR
    for rel in document.part.rels.values():
        if "image" in rel.reltype:
            try:
                img_bytes = rel.target_part.blob
                ocr_text, mean_conf = ocr_image_bytes(img_bytes, doc_id, "docx_image")
                if ocr_text.strip():
                    blocks.append(
                        Block(
                            doc_id=doc_id,
                            page=page_num,
                            text=ocr_text,
                            source="ocr",
                            ocr_conf=mean_conf,
                        )
                    )
            except Exception as exc:
                log.warning("docx_image_ocr_failed", doc_id=doc_id, error=str(exc))

    return blocks, {"page_count": 1, "has_ocr_pages": False}


import re as _re  # noqa: E402
_DOCX_HEADING_STYLES_RE = _re.compile(r"^Heading\s*\d+$", _re.IGNORECASE)


# ── XLSX parser ───────────────────────────────────────────────────────────────


def _parse_xlsx(raw: bytes, doc_id: str, policy: dict[str, Any]) -> tuple[list[Block], dict[str, Any]]:
    """Parse an XLSX using openpyxl."""
    try:
        import io

        import openpyxl  # type: ignore[import-untyped]
    except ImportError as exc:
        raise IngestionError("openpyxl is not installed") from exc

    _check_zip_bomb(raw, "xlsx", policy.get("limits", {}).get("zip_bomb_ratio", 50))

    wb = openpyxl.load_workbook(io.BytesIO(raw), data_only=True, read_only=True)
    blocks: list[Block] = []
    has_hidden_sheets = False

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        sheet_state = getattr(ws, "sheet_state", "visible")
        is_hidden = sheet_state != "visible"
        if is_hidden:
            has_hidden_sheets = True

        for row in ws.iter_rows():
            for cell in row:
                if cell.value is None:
                    continue
                cell_text = str(cell.value).strip()
                if not cell_text:
                    continue
                section_path = [sheet_name, f"{cell.column_letter}{cell.row}"]
                block = Block(
                    doc_id=doc_id,
                    page=0,
                    text=cell_text,
                    source="table_cell",
                    section_path=section_path,
                )
                if is_hidden:
                    # Tag hidden sheet blocks in text for downstream flagging
                    block = block.model_copy(
                        update={"text": f"[HIDDEN_SHEET] {cell_text}"}
                    )
                blocks.append(block)

    wb.close()
    return blocks, {"page_count": 1, "has_hidden_sheets": has_hidden_sheets}
